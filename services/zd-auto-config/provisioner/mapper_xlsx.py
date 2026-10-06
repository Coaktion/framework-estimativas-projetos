# -*- coding: utf-8 -*-
"""Mapper: planilha rev2 -> Blueprint.

A planilha CONTEM tudo. A decisao do que provisionar e feita no APP (nao no
_Controle). Por isso, aqui carregamos TODO o conteudo presente; o _Controle
guarda apenas os parametros (Identificador, Prefixo_TAG, Idioma da demo).
O Criar? por LINHA dentro de cada aba continua respeitado.
"""
from __future__ import annotations
import openpyxl
from blueprint import (Blueprint, Access0, Brand, Group, CustomField, FieldType, Lembrete,
                       FieldCondition, View, DefaultTrigger, Channel, Macro, SLA,
                       GuideSection, GuideArticle, Schedule, Holiday, Locale, Access,
                       CopilotSuggestion, AIAgentSuggestion)


def _rows(ws):
    headers = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        yield dict(zip(headers, r))


def _b(v) -> bool:
    return str(v).strip().upper() in ("S", "SIM", "TRUE", "1", "Y")


def load_blueprint(path: str) -> Blueprint:
    wb = openpyxl.load_workbook(path, data_only=True)
    S = wb.sheetnames

    # ---- _Controle: apenas parametros (col0 = rotulo, col1 = valor) ----
    params = {}
    for row in wb["_Controle"].iter_rows(min_row=2, values_only=True):
        label = (row[0] or "").strip() if row and row[0] else ""
        val = row[1] if len(row) > 1 else None
        if label and val not in (None, ""):
            params[label] = str(val).strip()
    idioma_raw = params.get("Idioma da demo", "pt-br")
    idioma = {"pt-br": Locale.pt, "en-us": Locale.en, "es": Locale.es}.get(idioma_raw, Locale.pt)

    # identificador = numero + emoji, SEM espaco entre eles (ex.: "056🛒", "11⚡")
    identificador = params.get("Identificador da demo", "").replace(" ", "")

    # ---- 0_Acesso (opcional) ----
    acc = next(_rows(wb["0_Acesso"]), {}) if "0_Acesso" in S else {}
    acesso = Access0(
        subdominio=str(acc.get("Zendesk (SEM https)", "")),
        admin_email=str(acc.get("Email do Admin", "")),
        auth_tipo="token" if "tok" in str(acc.get("Senha ou Token?", "token")).lower() else "senha",
        segredo=str(acc.get("Pwd/Token", "") or ""),
        config=str(acc.get("Config (simple/full)", "full")),
    )

    # ---- 0_Marca ----
    marca = None
    if "0_Marca" in S:
        m = next(_rows(wb["0_Marca"]), None)
        if m and m.get("Marca"):
            marca = Brand(criar=True, nome=str(m.get("Marca", "")),
                          subdominio=str(m.get("Subdominio", "") or ""),
                          tipo=str(m.get("Tipo (simple/full)", "full")))

    # ---- 1_Grupos ----
    grupos = []
    if "1_Grupos" in S:
        gs = sorted(_rows(wb["1_Grupos"]), key=lambda x: x.get("Ordem", 999) or 999)
        grupos = [Group(nome=str(g["Grupo"])) for g in gs if g.get("Grupo")]

    # ---- campos (ticket / usuario / org) ----
    def campos(sheet):
        out = []
        if sheet not in S:
            return out
        for row in _rows(wb[sheet]):
            if "Criar?" in row and not _b(row.get("Criar?")):
                continue
            out.append(CustomField(
                nome=str(row.get("Nome do Campo", "")),
                tipo=FieldType(str(row.get("Tipo", "text")).strip()),
                chave_ref=str(row.get("Chave_ref", "")).strip(),
                opcoes=str(row.get("Opcoes (se lista, por virgula)", row.get("Opcoes (se lista)", "")) or ""),
                visivel_portal=_b(row.get("Visivel portal?", "S")),
                nome_portal=(str(row.get("Nome no portal")) if row.get("Nome no portal") else None),
            ))
        return out

    campos_ticket = campos("2_Campos_Ticket")
    campos_usuario = campos("5_Campos_Usuario")
    campos_org = campos("6_Campos_Org")

    # ---- 2b_Condicionais ----
    condicionais = []
    if "2b_Condicionais" in S:
        split = lambda s: [x.strip() for x in str(s or "").split(",") if x.strip()]
        for row in _rows(wb["2b_Condicionais"]):
            if not _b(row.get("Criar?", "S")):
                continue
            condicionais.append(FieldCondition(
                campo_pai_ref=str(row.get("Campo Pai (Chave_ref)", "")).strip(),
                quando_valor=str(row.get("Quando valor for", "")).strip(),
                mostrar_refs=split(row.get("Mostrar campos (Chave_ref)")),
                obrigatorios_refs=split(row.get("Tornar obrigatorios (Chave_ref)")),
                publico=str(row.get("Publico", "ambos")).strip()))

    # ---- 3_Views ----
    views = []
    if "3_Views" in S:
        for row in _rows(wb["3_Views"]):
            if not _b(row.get("Criar?", "S")):
                continue
            views.append(View(categoria=str(row.get("Grupo (categoria visual)", "")),
                              nome=str(row.get("Nome da View", "")),
                              criterio=str(row.get("Definicao / criterio", ""))))

    # ---- 4_Gatilhos_Padrao ----
    gatilhos = []
    if "4_Gatilhos_Padrao" in S:
        for row in _rows(wb["4_Gatilhos_Padrao"]):
            if not _b(row.get("Criar?", "S")):
                continue
            gatilhos.append(DefaultTrigger(
                nome=str(row.get("Nome do Gatilho", "")),
                canal=Channel(str(row.get("Canal", "web")).strip()),
                copilot=_b(row.get("Copilot?", "N"))))

    # ---- 8_Macros ----
    macros = []
    if "8_Macros" in S:
        for row in _rows(wb["8_Macros"]):
            if not _b(row.get("Criar?", "S")):
                continue
            macros.append(Macro(
                nome=str(row.get("Nome", "")),
                mensagem=str(row.get("Mensagem", "")),
                status_ticket=(str(row["Status ticket"]) if row.get("Status ticket") else None),
                prioridade=(str(row["Prioridade"]) if row.get("Prioridade") else None),
                publico=_b(row.get("Publico?", "S")),
                grupo_visualizador=(str(row["Grupo visualizador"]) if row.get("Grupo visualizador") else None),
                grupo_destino=(str(row["Grupo destino"]) if row.get("Grupo destino") else None)))

    # ---- 9_SLA ----
    sla = None
    if "9_SLA" in S:
        s = next(_rows(wb["9_SLA"]), None)
        if s:
            keys = ["urgent", "high", "normal", "low"]
            fr = ["Urgente 1a Resp (min)", "Alta 1a Resp", "Normal 1a Resp", "Baixa 1a Resp"]
            rs = ["Urgente Resol", "Alta Resol", "Normal Resol", "Baixa Resol"]
            sla = SLA(criar=True,
                      first_reply={k: int(s[c]) for k, c in zip(keys, fr) if s.get(c) is not None},
                      resolution={k: int(s[c]) for k, c in zip(keys, rs) if s.get(c) is not None})

    # ---- Guide ----
    secoes, artigos = [], []
    if "10_Guide_Secoes" in S:
        for row in _rows(wb["10_Guide_Secoes"]):
            if _b(row.get("Criar?", "S")):
                secoes.append(GuideSection(titulo=str(row.get("Titulo", ""))))
    if "10_Guide_Artigos" in S:
        for row in _rows(wb["10_Guide_Artigos"]):
            if _b(row.get("Criar?", "S")):
                artigos.append(GuideArticle(
                    secao_titulo=str(row.get("Secao (Titulo)", "")),
                    titulo=str(row.get("Titulo", "")),
                    corpo=str(row.get("Corpo", "")),
                    acesso=Access(str(row.get("Acesso", "aberto")).strip())))

    # ---- Programacao + Feriados ----
    programacao, feriados = None, []
    if "12_Programacao" in S:
        p = next(_rows(wb["12_Programacao"]), None)
        if p and p.get("Nome"):
            programacao = Schedule(criar=True, nome=str(p.get("Nome", "Comercial")),
                                   dia_inicio=str(p.get("Dia inicio", "seg")),
                                   dia_fim=str(p.get("Dia fim", "sex")),
                                   hora_inicio=int(p.get("Hora inicio", 8)),
                                   hora_fim=int(p.get("Hora fim", 18)),
                                   feriados_modo=(str(p.get("Feriados (normal/prolongado)")) or None))
    if "12_Feriados" in S:
        for row in _rows(wb["12_Feriados"]):
            if row.get("Nome"):
                feriados.append(Holiday(nome=str(row.get("Nome", "")), ano=int(row.get("Ano", 2026)),
                                        mes=int(row.get("Mes", 1)), dia_inicio=int(row.get("Dia inicio", 1)),
                                        dia_fim=int(row.get("Dia fim", 1))))

    # ---- _Guia_Apresentacao (artigo interno) ----
    guia_apres = None
    if "_Guia_Apresentacao" in S:
        vals = {}
        for row in wb["_Guia_Apresentacao"].iter_rows(values_only=True):
            if row and row[0]:
                vals[str(row[0]).strip().lower()] = row[1]
        titulo = vals.get("titulo") or vals.get("titulo (agente)")
        corpo = vals.get("corpo")
        if titulo and corpo and "reservado" not in str(corpo).lower():
            guia_apres = GuideArticle(secao_titulo="Guia de Apresentacao (interno)",
                                      titulo=str(titulo), corpo=str(corpo), acesso=Access.agente)

    # ---- _Lembretes ----
    lembretes = []
    if "_Lembretes" in S:
        for row in _rows(wb["_Lembretes"]):
            t = row.get("Titulo") or row.get("Título")
            if t:
                lembretes.append(Lembrete(titulo=str(t), detalhe=str(row.get("Detalhe", "") or "")))

    # ---- _Copilot (sugestoes de procedure; informativo, sem API) ----
    copilots = []
    if "_Copilot" in S:
        for row in _rows(wb["_Copilot"]):
            t = row.get("Titulo") or row.get("Título")
            if t:
                copilots.append(CopilotSuggestion(
                    titulo=str(t),
                    quando_utilizar=str(row.get("Quando utilizar", "") or ""),
                    corpo=str(row.get("Corpo", "") or "")))

    # ---- _AI_Agents (sugestoes de fluxo; informativo, sem API) ----
    ai_agents = []
    if "_AI_Agents" in S:
        for row in _rows(wb["_AI_Agents"]):
            n = row.get("Nome")
            if n:
                ai_agents.append(AIAgentSuggestion(
                    nome=str(n),
                    descricao=str(row.get("Descricao", row.get("Descrição", "")) or ""),
                    prompt=str(row.get("Prompt", "") or "")))

    return Blueprint(
        cliente=(marca.nome if marca else "Demo"),
        tag_prefix=params.get("Prefixo_TAG (global desta demo)", "demo"),
        identificador=identificador,
        idioma=idioma,
        acesso=acesso, marca=marca, grupos=grupos,
        campos_ticket=campos_ticket, condicionais=condicionais,
        views=views, gatilhos_padrao=gatilhos,
        campos_usuario=campos_usuario, campos_org=campos_org,
        macros=macros, sla=sla, guide_secoes=secoes, guide_artigos=artigos,
        programacao=programacao, feriados=feriados,
        guia_apresentacao=guia_apres, lembretes=lembretes,
        copilots=copilots, ai_agents=ai_agents,
    )
